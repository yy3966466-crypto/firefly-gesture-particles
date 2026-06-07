import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document('我的论文_v6.docx')

def replace_para_text(paragraph, new_text):
    """Replace all runs in a paragraph with a single run containing new_text,
    preserving the first run's formatting."""
    p = paragraph._p
    pPr = p.find(qn('w:pPr'))

    # Get formatting from first run
    first_rPr = None
    first_run = p.find(qn('w:r'))
    if first_run is not None:
        first_rPr = first_run.find(qn('w:rPr'))

    # Remove all runs
    for r in p.findall(qn('w:r')):
        p.remove(r)

    # Create new run with preserved formatting
    new_r = OxmlElement('w:r')
    if first_rPr is not None:
        new_r.append(first_rPr.__deepcopy__({}))

    new_t = OxmlElement('w:t')
    new_t.text = new_text
    new_t.set(qn('xml:space'), 'preserve')
    new_r.append(new_t)

    # Insert after pPr or at beginning
    if pPr is not None:
        p.insert(p.index(pPr) + 1, new_r)
    else:
        p.insert(0, new_r)


# ---- Para 26: intro sentence ----
para26_text = '本研究以MATLAB为核心平台，围绕多生理参数监护系统的算法设计与系统集成展开了研究，主要研究内容包括以下四个方面：'
replace_para_text(doc.paragraphs[26], para26_text)
print('Para 26 updated')

# ---- Para 27: (1) 生理信号预处理 ----
para27_text = '（1）生理信号预处理。基于MIT-BIH心律失常数据库和PhysioNet呼吸信号数据库的真实数据，结合了实验室实测的生理信号，构建了生理信号数据集。设计了滑动中值滤波基线校准、频域带通滤波和移动平均平滑等预处理算法，有效抑制了工频干扰、肌电干扰和基线漂移等噪声，为后续的特征提取提供了可靠的信号基础。'
replace_para_text(doc.paragraphs[27], para27_text)
print('Para 27 updated')

# ---- Para 28: (2) 关键信号处理算法 ----
para28_text = '（2）关键信号处理算法设计与优化。针对心电信号设计了基于频域带通滤波与自适应差分阈值的R波检测方法；针对呼吸信号设计了基于LMS自适应对消的呼吸率提取方案；针对体温信号设计了数字滤波与线性校准的体温处理算法。在MATLAB环境下完成了各算法的仿真实现与参数优化。'
replace_para_text(doc.paragraphs[28], para28_text)
print('Para 28 updated')

# ---- Para 29: (3) 综合监护平台 ----
para29_text = '（3）综合监护平台构建。基于MATLAB GUI搭建了集实时波形显示、心率呼吸率体温等关键参数计算、数据存储和异常报警于一体的综合监护平台，实现了心电、呼吸、体温信号的同步监测与多参数协同处理。'
replace_para_text(doc.paragraphs[29], para29_text)
print('Para 29 updated')

# ---- Para 30: (4) 系统性能验证 ----
para30_text = '（4）系统性能验证。从准确性、稳定性和抗干扰能力三个维度对系统进行了实验验证，通过信噪比（SNR）、均方误差（MSE）、检测准确率等量化指标评估了算法的有效性。'
replace_para_text(doc.paragraphs[30], para30_text)
print('Para 30 updated')

# ---- Para 31: 技术路线 ----
para31_text = '本研究采取的技术路线为：文献调研与理论学习→数据集构建与预处理→算法设计与MATLAB仿真→参数优化与测试→综合监护平台构建→系统性能实验验证→结果分析与迭代完善。'
replace_para_text(doc.paragraphs[31], para31_text)
print('Para 31 updated')

# ---- Save ----
doc.save('我的论文_v6_updated.docx')
print('\nDone → 我的论文_v6_updated.docx')
