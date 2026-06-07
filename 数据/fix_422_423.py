"""Fix 4.2.2 RR formula and add 4.2.3 body content."""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy

doc = Document('d:/bishe/我的论文.docx')
MATH_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'

def make_rPr():
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Cambria Math')
    rFonts.set(qn('w:hAnsi'), 'Cambria Math')
    rPr.append(rFonts)
    return rPr

def make_ctrlPr():
    ctrlPr = OxmlElement('m:ctrlPr')
    ctrlPr.append(make_rPr())
    return ctrlPr

def make_mr(text):
    r = OxmlElement('m:r')
    t = OxmlElement('m:t'); t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    return r

def make_text_run(text):
    r = OxmlElement('w:r')
    t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve')
    r.append(t)
    return r

def make_inline_om(text):
    """Simple inline OMML variable."""
    om = OxmlElement('m:oMath')
    om.append(make_mr(text))
    return om

def make_subscript_om(base, sub_text):
    """Inline OMML with m:sSub."""
    om = OxmlElement('m:oMath')
    sSub = OxmlElement('m:sSub')
    sSubPr = OxmlElement('m:sSubPr')
    sSubPr.append(make_ctrlPr())
    sSub.append(sSubPr)

    e = OxmlElement('m:e')
    e.append(make_mr(base))
    e.append(make_ctrlPr())
    sSub.append(e)

    sub_el = OxmlElement('m:sub')
    sub_el.append(make_mr(sub_text))
    sub_el.append(make_ctrlPr())
    sSub.append(sub_el)

    om.append(sSub)
    return om

# ═══════════════════════════════════════════════════════════════
# FIX 1: Replace broken RR formula in para 136
# ═══════════════════════════════════════════════════════════════
p136 = doc.paragraphs[136]

# Find the last OMML element (broken RR formula)
omaths = p136._element.findall(f'{{{MATH_NS}}}oMath')
if omaths:
    last_om = omaths[-1]
    # Get the parent's index of this element
    parent = last_om.getparent()
    # Check if it's really the broken RR formula
    texts = []
    for t_el in last_om.iter(f'{{{MATH_NS}}}t'):
        if t_el.text:
            texts.append(t_el.text)
    print(f'Last OMML content: \"{"".join(texts)}\"')

    # Remove the broken element
    parent.remove(last_om)
    print('Removed broken RR formula')

    # Create correct RR formula
    new_rr = OxmlElement('m:oMath')

    # RR=
    new_rr.append(make_mr('RR='))

    # Fraction 60/T_resp
    f = OxmlElement('m:f')
    fPr = OxmlElement('m:fPr')
    fPr.append(make_ctrlPr())
    f.append(fPr)

    # Numerator: 60
    num = OxmlElement('m:num')
    num.append(make_mr('60'))
    num.append(make_ctrlPr())
    f.append(num)

    # Denominator: T_resp (subscript)
    den = OxmlElement('m:den')
    sSub = OxmlElement('m:sSub')
    sSubPr = OxmlElement('m:sSubPr')
    sSubPr.append(make_ctrlPr())
    sSub.append(sSubPr)

    # Base: T
    den_e = OxmlElement('m:e')
    den_e.append(make_mr('T'))
    den_e.append(make_ctrlPr())
    sSub.append(den_e)

    # Subscript: resp
    den_sub = OxmlElement('m:sub')
    den_sub.append(make_mr('resp'))
    den_sub.append(make_ctrlPr())
    sSub.append(den_sub)

    den.append(sSub)
    den.append(make_ctrlPr())
    f.append(den)

    new_rr.append(f)

    # Append the new correct formula
    parent.append(new_rr)
    print('Added corrected RR formula')

    # Verify
    new_texts = []
    for t_el in new_rr.iter(f'{{{MATH_NS}}}t'):
        if t_el.text:
            new_texts.append(t_el.text)
    print(f'New content: \"{"".join(new_texts)}\"')
else:
    print('No OMML elements found in para 136!')

# ═══════════════════════════════════════════════════════════════
# FIX 2: Add body text to 4.2.3 section
# ═══════════════════════════════════════════════════════════════
# Current state:
#   [137]: 4.2.3 header
#   [138]: 图4.5
# Need to insert text between them

p137 = doc.paragraphs[137]  # 4.2.3 header
p138 = doc.paragraphs[138]  # 图4.5

print(f'\nPara 137 (header): {p137.text.strip()[:80]}')
print(f'Para 138 (after): {p138.text.strip()[:80]}')

# Build the body text paragraph with inline OMML for variables
body_p = OxmlElement('w:p')

# Paragraph properties: first line indent, spacing
pPr = OxmlElement('w:pPr')
ind = OxmlElement('w:ind')
ind.set(qn('w:firstLine'), '480')
pPr.append(ind)
spacing = OxmlElement('w:spacing')
spacing.set(qn('w:after'), '120')
pPr.append(spacing)
body_p.append(pPr)

# Text content with OMML variables
# Section describes parameter selection and evaluation results
body_segments = [
    ('text', '为了验证LMS自适应线增强器对呼吸信号的处理效果，利用PhysioNet呼吸信号数据库中的数据进行测试。自适应线增强器的参数依据信号采样率自适应设定：滤波器阶数 '),
    ('math', 'order = max(8, min(40, round(0.40 × f_s)))'),
    ('text', '，去相关延迟 '),
    ('math', 'delay = max(order + 2, round(0.25 × f_s))'),
    ('text', '，步长因子 '),
    ('math', 'μ = 0.015 / (order × var(signal))'),
    ('text', '。上述参数保证了滤波器在收敛速度与稳态误差之间取得合理平衡。'),
    ('text', '经过ALE增强处理后，呼吸信号的信噪比得到显著提升。对增强信号进一步进行移动平均平滑处理（窗口长度 '),
    ('math', '0.95 × f_s'),
    ('text', '），消除残余的高频波动。滤波前后的呼吸信号波形对比如图4.5所示，可以看出ALE有效抑制了心电干扰和运动伪差，保留了呼吸信号的周期性特征。'),
]

for seg_type, seg_content in body_segments:
    if seg_type == 'text':
        body_p.append(make_text_run(seg_content))
    elif seg_type == 'math':
        body_p.append(make_inline_om(seg_content))

# Insert the body paragraph after the 4.2.3 header
# addnext inserts as the immediate next sibling
p137._element.addnext(body_p)
print('Inserted 4.2.3 body text after header')

# ═══════════════════════════════════════════════════════════════
# Verify
# ═══════════════════════════════════════════════════════════════
print('\n=== Verification ===')

# Check 4.2.2 para 136 - RR formula
p136 = doc.paragraphs[136]
omaths = p136._element.findall(f'{{{MATH_NS}}}oMath')
print(f'Para 136 OMML count: {len(omaths)}')
last_texts = []
for t_el in omaths[-1].iter(f'{{{MATH_NS}}}t'):
    if t_el.text:
        last_texts.append(t_el.text)
print(f'Last OMML (RR formula): \"{"".join(last_texts)}\"')
# Check for num/den in fraction
fracs = omaths[-1].findall(f'{{{MATH_NS}}}f')
for f in fracs:
    num = f.find(f'{{{MATH_NS}}}num')
    den = f.find(f'{{{MATH_NS}}}den')
    print(f'  Has num: {num is not None}, Has den: {den is not None}')

# Check 4.2.3 area
print('\n4.2.3 area:')
for i in range(137, 142):
    if i < len(doc.paragraphs):
        t = doc.paragraphs[i].text.strip()[:100]
        print(f'  [{i}] {t}')

doc.save('d:/bishe/我的论文.docx')
print('\nSaved successfully')
