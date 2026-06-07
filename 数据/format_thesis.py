import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document('我的论文_v6_revised.docx')

# ================================================================
# 1) Intro text update (paras 26-31)
# ================================================================
def replace_para_text(paragraph, new_text):
    p = paragraph._p
    pPr = p.find(qn('w:pPr'))
    first_rPr = None
    first_run = p.find(qn('w:r'))
    if first_run is not None:
        first_rPr = first_run.find(qn('w:rPr'))
    for r in p.findall(qn('w:r')):
        p.remove(r)
    new_r = OxmlElement('w:r')
    if first_rPr is not None:
        new_r.append(first_rPr.__deepcopy__({}))
    new_t = OxmlElement('w:t')
    new_t.text = new_text
    new_t.set(qn('xml:space'), 'preserve')
    new_r.append(new_t)
    if pPr is not None:
        p.insert(p.index(pPr) + 1, new_r)
    else:
        p.insert(0, new_r)

texts = {
    26: '本研究以MATLAB为核心平台，围绕多生理参数监护系统的算法设计与系统集成展开了研究，主要研究内容包括以下四个方面：',
    27: '（1）生理信号预处理。基于MIT-BIH心律失常数据库和PhysioNet呼吸信号数据库的真实数据，结合了实验室实测的生理信号，构建了生理信号数据集。设计了滑动中值滤波基线校准、频域带通滤波和移动平均平滑等预处理算法，有效抑制了工频干扰、肌电干扰和基线漂移等噪声，为后续的特征提取提供了可靠的信号基础。',
    28: '（2）关键信号处理算法设计与优化。针对心电信号设计了基于频域带通滤波与自适应差分阈值的R波检测方法；针对呼吸信号设计了基于LMS自适应对消的呼吸率提取方案；针对体温信号设计了数字滤波与线性校准的体温处理算法。在MATLAB环境下完成了各算法的仿真实现与参数优化。',
    29: '（3）综合监护平台构建。基于MATLAB GUI搭建了集实时波形显示、心率呼吸率体温等关键参数计算、数据存储和异常报警于一体的综合监护平台，实现了心电、呼吸、体温信号的同步监测与多参数协同处理。',
    30: '（4）系统性能验证。从准确性、稳定性和抗干扰能力三个维度对系统进行了实验验证，通过信噪比（SNR）、均方误差（MSE）、检测准确率等量化指标评估了算法的有效性。',
    31: '本研究采取的技术路线为：文献调研与理论学习→数据集构建与预处理→算法设计与MATLAB仿真→参数优化与测试→综合监护平台构建→系统性能实验验证→结果分析与迭代完善。',
}
for idx, text in texts.items():
    replace_para_text(doc.paragraphs[idx], text)
    print(f'[Para {idx}] intro text updated')

# ================================================================
# 2) Figure renumbering (Chapter 3)
# ================================================================
fig_fixes = {
    79: ('图3.2', '图3.1'),
    88: ('图3.3', '图3.2'),
    98: ('图3.4', '图3.3'),
    101: ('图3.1', '图3.4'),
}
for idx, (old, new) in fig_fixes.items():
    for run in doc.paragraphs[idx].runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
    print(f'[Para {idx}] {old} → {new}')

for run in doc.paragraphs[99].runs:
    if '图3.1' in run.text:
        run.text = run.text.replace('图3.1', '图3.4')
        print(f'[Para 99] 图3.1 → 图3.4')
        break

# ================================================================
# 3) Three-line tables
# ================================================================
def make_three_line(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    old_borders = tblPr.find(qn('w:tblBorders'))
    if old_borders is not None:
        tblPr.remove(old_borders)
    borders = OxmlElement('w:tblBorders')
    for side, val, sz in [('top', 'single', 12), ('bottom', 'single', 6),
                           ('left', 'none', 0), ('right', 'none', 0),
                           ('insideH', 'none', 0), ('insideV', 'none', 0)]:
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), val)
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tblPr.append(borders)
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.tcPr
            if tcPr is not None:
                tc_borders = tcPr.find(qn('w:tcBorders'))
                if tc_borders is not None:
                    tcPr.remove(tc_borders)
    for cell in table.rows[0].cells:
        tcPr = cell._tc.tcPr
        if tcPr is None:
            tcPr = OxmlElement('w:tcPr')
            cell._tc.insert(0, tcPr)
        tc_borders = OxmlElement('w:tcBorders')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '0')
        bottom.set(qn('w:color'), '000000')
        tc_borders.append(bottom)
        tcPr.append(tc_borders)

for i, table in enumerate(doc.tables):
    make_three_line(table)
    print(f'[Table {i}] three-line')

# ================================================================
# 4) Save
# ================================================================
output_path = '我的论文_v6_final_formatted.docx'
doc.save(output_path)
print(f'\nDone → {output_path}')
