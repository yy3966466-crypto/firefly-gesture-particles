"""Insert respiration comparison figure into thesis at 图4.5."""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Cm

doc = Document('d:/bishe/我的论文.docx')

# ── 1. Add the picture at the end of the document ──
# python-docx handles image embedding automatically
pic_para = doc.add_paragraph()
pic_run = pic_para.add_run()
pic_run.add_picture('d:/bishe/fig_respiration_comparison.png', width=Cm(14))

# Set paragraph alignment to center
from docx.enum.text import WD_ALIGN_PARAGRAPH
pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ── 2. Move the picture paragraph to the right position ──
# Current structure:
#   [137]: 4.2.3 header
#   [138]: body text
#   [139]: 图4.5 caption
# Need to insert the image between [138] and [139]

body = doc.element.body
# Find the picture paragraph we just added (it's the last paragraph)
all_paras = list(body.iterchildren(qn('w:p')))
new_pic_elem = all_paras[-1]  # The picture paragraph

# Find para 139 element (caption)
p139 = doc.paragraphs[139]
caption_elem = p139._element

# Insert the picture before the caption
caption_elem.addprevious(new_pic_elem)

print(f'Inserted image paragraph before para 139 (caption)')

# ── 3. Verify ──
print('\nVerification:')
for i in range(137, 143):
    if i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        t = p.text.strip()
        has_img = False
        for child in p._element:
            if 'drawing' in child.tag:
                has_img = True
                break
            for sub in child:
                if 'drawing' in sub.tag:
                    has_img = True
                    break
        print(f'  [{i}] img={has_img} \"{t[:80]}\"')

doc.save('d:/bishe/我的论文.docx')
print('\nDone')
