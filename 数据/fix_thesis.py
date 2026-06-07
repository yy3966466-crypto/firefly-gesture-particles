import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

MATH_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
etree.register_namespace('m', MATH_NS)

def m(tag):
    return OxmlElement(f'm:{tag}')

def mr(text):
    r = m('r')
    t = m('t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    return r

def set_equation(paragraph, text, num_text=None):
    """Replace paragraph with OMML equation containing plain text."""
    p = paragraph._p
    pPr = p.find(qn('w:pPr'))
    for child in list(p):
        if child.tag != qn('w:pPr'):
            p.remove(child)

    omath = m('oMath')
    omath.append(mr(text))
    omathPara = m('oMathPara')
    omathPara.append(omath)

    if pPr is not None:
        p.insert(p.index(pPr) + 1, omathPara)
    else:
        p.append(omathPara)

    if num_text:
        tab_run = OxmlElement('w:r')
        tab = OxmlElement('w:tab')
        tab_run.append(tab)
        p.append(tab_run)
        num_run = OxmlElement('w:r')
        num_t_w = OxmlElement('w:t')
        num_t_w.text = f'({num_text})'
        num_t_w.set(qn('xml:space'), 'preserve')
        num_run.append(num_t_w)
        p.append(num_run)

def make_three_line(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)

    old_borders = tblPr.find(qn('w:tblBorders'))
    if old_borders is not None:
        tblPr.remove(old_borders)

    borders = OxmlElement('w:tblBorders')
    for side, val, sz in [('top', 'single', 12), ('bottom', 'single', 6),
                           ('left', 'none', 0), ('right', 'none', 0),
                           ('insideH', 'none', 0), ('insideV', 'none', 0)]:
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), val)
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tblPr.append(borders)

    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.tcPr
            if tcPr is not None:
                tc_borders = tcPr.find(qn('w:tcBorders'))
                if tc_borders is not None:
                    tcPr.remove(tc_borders)

    for cell in table.rows[0].cells:
        tcPr = cell._tc.tcPr
        if tcPr is None:
            tcPr = OxmlElement('w:tcPr')
            cell._tc.insert(0, tcPr)
        tc_borders = OxmlElement('w:tcBorders')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '0')
        bottom.set(qn('w:color'), '000000')
        tc_borders.append(bottom)
        tcPr.append(tc_borders)

# ================================================================
doc = Document('我的论文_v6_updated.docx')

# --- Figure renumbering (Chapter 3) ---
fig_fixes = {
    79: ('图3.2', '图3.1'),
    88: ('图3.3', '图3.2'),
    98: ('图3.4', '图3.3'),
    101: ('图3.1', '图3.4'),
}
for idx, (old, new) in fig_fixes.items():
    for run in doc.paragraphs[idx].runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
    print(f'[Para {idx}] {old} → {new}')

for run in doc.paragraphs[99].runs:
    if '图3.1' in run.text:
        run.text = run.text.replace('图3.1', '图3.4')
        print(f'[Para 99] 图3.1 → 图3.4')
        break

# --- OMML Equations ---
eq_defs = [
    (49, "x'(n) = (1/N) ∑_{k=0}^{N-1} X(k)·M(k)·e^{j2πkn/N}", '2-1'),
    (50, 'X(k) = ∑_{n=0}^{N-1} x(n)·e^{−j2πkn/N}', '2-2'),
    (111, 'threshold(n) = mean(envelope(n)) + 1.40 × std(envelope(n))', '4-1'),
    (127, 'y(n) = w^{T}(n)·x(n−d)', '4-2'),
    (126, 'e(n) = d(n) − y(n)', '4-3'),
    (125, 'w(n+1) = w(n) + 2μ·e(n)·x(n−d)', '4-4'),
    (167, 'CV = σ / μ', '5-1'),
    (197, 'SNR(dB) = 10 × log₁₀(var(x_filtered) / var(x_raw − x_filtered))', '6-1'),
]

for idx, text, num in eq_defs:
    set_equation(doc.paragraphs[idx], text, num)
    print(f'[Para {idx}] Equation ({num})')

# --- Three-line tables ---
for i, table in enumerate(doc.tables):
    make_three_line(table)
    print(f'[Table {i}] three-line')

# --- Save ---
doc.save('我的论文_v6_formatted.docx')
print('\nDone!')
