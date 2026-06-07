"""Verify final state of all fixes."""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

doc = Document('d:/bishe/我的论文.docx')
MATH_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'

print('=== Para 212 (fixed) ===')
print(doc.paragraphs[212].text.strip())

print('\n=== Para 201 OMML count ===')
oms = doc.paragraphs[201]._element.findall(f'{{{MATH_NS}}}oMath')
print(f'OMML count: {len(oms)}')
for j, om in enumerate(oms):
    texts = []
    for t in om.iter(f'{{{MATH_NS}}}t'):
        if t.text: texts.append(t.text)
    print(f'  [{j}] {"".join(texts)}')

print('\n=== Para 202 ===')
print(doc.paragraphs[202].text.strip()[:120])
oms2 = doc.paragraphs[202]._element.findall(f'{{{MATH_NS}}}oMath')
print(f'OMML count: {len(oms2)}')

print('\n=== Table 3 Row 5 (233) ===')
row = doc.tables[3].rows[5]
print(f'TP: {row.cells[3].text}')

print('\n=== Table 6 HR/RR rows ===')
t6 = doc.tables[6]
for ri in [1, 2]:
    cells = [t6.rows[ri].cells[ci].text for ci in range(4)]
    print(f'Row {ri}: {cells}')

print('\nDone.')
