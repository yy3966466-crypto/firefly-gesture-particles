"""Re-insert Figure 4.5 image and verify final 4.2 structure."""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document('d:/bishe/我的论文.docx')
body = doc.element.body

print('=== Current structure ===')
for i in range(135, 145):
    if i < len(doc.paragraphs):
        t = doc.paragraphs[i].text.strip()[:90]
        has_img = False
        for child in doc.paragraphs[i]._element:
            if 'drawing' in child.tag: has_img = True; break
            for sub in child:
                if 'drawing' in sub.tag: has_img = True; break
        print(f'[{i}] img={has_img} \"{t}\"')

# Insert image between para 139 (4.2.3 body) and para 140 (图4.5 caption)
# Strategy: add picture to end, then move before caption
pic_para = doc.add_paragraph()
pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
pic_run = pic_para.add_run()
pic_run.add_picture('d:/bishe/fig_respiration_comparison.png', width=Cm(14))

# Move before caption
all_paras = list(body.iterchildren(qn('w:p')))
new_pic_elem = all_paras[-1]

p140 = doc.paragraphs[140]  # 图4.5 caption
p140._element.addprevious(new_pic_elem)

print('\nInserted image before para 140 (caption)')

# Verify final structure
print('\n=== Final structure ===')
for i in range(135, 150):
    if i < len(doc.paragraphs):
        t = doc.paragraphs[i].text.strip()[:90]
        has_img = False
        for child in doc.paragraphs[i]._element:
            if 'drawing' in child.tag: has_img = True; break
            for sub in child:
                if 'drawing' in sub.tag: has_img = True; break
        print(f'[{i}] img={has_img} \"{t}\"')

doc.save('d:/bishe/我的论文.docx')
print('\nDone')
