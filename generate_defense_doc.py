#!/usr/bin/env python3
"""Convert 答辩文稿及问答准备.md to formatted Word document."""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
import re

doc = Document()

# ---- Page setup ----
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# ---- Style definitions ----
style = doc.styles['Normal']
font = style.font
font.name = '宋体'
font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# Heading styles
for level, (size, color_hex) in enumerate([
    (22, '155BC3'),  # Heading 1
    (16, '155BC3'),  # Heading 2
    (14, '3485F5'),  # Heading 3
    (12, '1A1A2E'),  # Heading 4
], start=1):
    h_style = doc.styles[f'Heading {level}']
    h_font = h_style.font
    h_font.name = '黑体'
    h_font.size = Pt(size)
    h_font.bold = True
    h_font.color.rgb = RGBColor.from_string(color_hex)
    h_style.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

BLUE = RGBColor(0x15, 0x5B, 0xC3)
DARK = RGBColor(0x1A, 0x1A, 0x2E)
GRAY = RGBColor(0x5A, 0x5F, 0x72)


def add_para(text, font_size=12, bold=False, color=DARK, font_name='宋体',
             alignment=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, space_before=0,
             first_line_indent=None):
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if first_line_indent:
        p.paragraph_format.first_line_indent = Cm(first_line_indent)
    run = p.add_run(text)
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    return p


def add_separator():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run('─' * 60)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0xDE, 0xE2, 0xEC)
    run.font.name = '宋体'


def parse_and_add(text):
    """Parse markdown text and add to document."""
    lines = text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]

        # Skip empty lines
        if not line.strip():
            i += 1
            continue

        # H1
        if line.startswith('# ') and not line.startswith('## '):
            doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue

        # H2
        if line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue

        # H3
        if line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue

        # H4
        if line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=4)
            i += 1
            continue

        # Separator
        if line.strip().startswith('---'):
            add_separator()
            i += 1
            continue

        # Bold text block (section headers like **xxx**)
        bold_match = re.match(r'^\*\*(.+)\*\*$', line.strip())
        if bold_match:
            add_para(bold_match.group(1), font_size=13, bold=True, color=BLUE,
                     space_before=8, space_after=4)
            i += 1
            continue

        # Bullet points
        if re.match(r'^[-•]\s+', line.strip()):
            text = re.sub(r'^[-•]\s+', '', line.strip())
            # Handle bold within bullet
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            p = doc.add_paragraph(style='List Bullet')
            p.clear()
            run = p.add_run(text)
            run.font.size = Pt(12)
            run.font.name = '宋体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            i += 1
            continue

        # Numbered items (Q1:, Q2:, etc.)
        q_match = re.match(r'^(###\s+)?(Q\d+)：(.+)', line.strip())
        if q_match:
            q_num = q_match.group(2)
            q_text = q_match.group(3)
            add_para(f'{q_num}：{q_text}', font_size=13, bold=True, color=BLUE,
                     space_before=10, space_after=2)
            i += 1
            continue

        # "答：" line
        if line.strip().startswith('**答：**') or line.strip().startswith('答：'):
            text = re.sub(r'^\*\*答：\*\*\s*', '', line.strip())
            text = re.sub(r'^答：\s*', '', text)
            add_para(f'答：{text}', font_size=12, bold=False, color=DARK,
                     first_line_indent=0.74, space_after=4)
            i += 1
            continue

        # Numbered list items like "第一，xxx"
        if re.match(r'^第[一二三四五六七八九十]', line.strip()):
            text = line.strip()
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            add_para(text, font_size=12, color=DARK, first_line_indent=0.74,
                     space_after=3)
            i += 1
            continue

        # Table rows (markdown tables)
        if '|' in line and line.strip().startswith('|'):
            i += 1
            continue

        # Regular paragraph
        text = line.strip()
        # Handle inline bold
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        # Handle inline code
        text = re.sub(r'`([^`]+)`', r'\1', text)
        if text:
            add_para(text, font_size=12, color=DARK, first_line_indent=0.74,
                     space_after=4)
        i += 1


# ---- Read markdown and parse ----
with open('d:/bishe/答辩文稿及问答准备.md', 'r', encoding='utf-8') as f:
    content = f.read()

parse_and_add(content)

# ---- Save ----
output_path = 'd:/bishe/答辩文稿及问答准备.docx'
doc.save(output_path)
print(f'Saved: {output_path}')
