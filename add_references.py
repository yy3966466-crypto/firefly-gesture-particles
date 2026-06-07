#!/usr/bin/env python3
"""Add missing references [4][9][14][18][21] using paragraph indices, marked in RED."""

from docx import Document
from docx.shared import Pt, RGBColor

RED = RGBColor(0xFF, 0x00, 0x00)
SRC = r"d:\bishe\多生理参数监护系统及关键信号处理算法.docx"

doc = Document(SRC)

def add_citation(para, cite_text, font_size=10.5):
    """Add cite_text as a RED run at end of paragraph."""
    # Trim trailing whitespace from last run
    if para.runs:
        last = para.runs[-1]
        last.text = last.text.rstrip()
    new_run = para.add_run(cite_text)
    new_run.font.size = Pt(font_size)
    new_run.font.name = '宋体'
    new_run.font.color.rgb = RED
    return new_run

results = []

# ================================================================
# [4] 王健琪 — LMS算法消噪应用
# ================================================================
# IDX=94: Section 2.2.3 LMS原理 (already has [15][16])
p = doc.paragraphs[94]
if "[4]" not in p.text:
    add_citation(p, "[4]")
    results.append(f"P94 (2.2.3 LMS原理) +[4]")

# IDX=154: Section 4.2.1 LMS ALE设计 (already has [15][16])
p = doc.paragraphs[154]
if "[4]" not in p.text:
    add_citation(p, "[4]")
    results.append(f"P154 (4.2.1 LMS ALE) +[4]")

# ================================================================
# [9] Arcos-Santiago — 包络解调应用于阻抗呼吸描记
# ================================================================
# IDX=72: Section 2.1.2 呼吸信号检测原理 (already has [15])
p = doc.paragraphs[72]
if "[9]" not in p.text:
    add_citation(p, "[9]")
    results.append(f"P72 (2.1.2 呼吸检测原理) +[9]")

# IDX=166: Section 4.2.2 呼吸峰值检测
p = doc.paragraphs[166]
if "[9]" not in p.text:
    add_citation(p, "[9]")
    results.append(f"P166 (4.2.2 呼吸峰检测) +[9]")

# ================================================================
# [14] 王蔷薇 — 基于提升小波的心电R波检测
# ================================================================
# IDX=52: Section 1.2.2 ECG算法研究现状 (already has [11][12][13])
p = doc.paragraphs[52]
if "[14]" not in p.text:
    add_citation(p, "[14]")
    results.append(f"P52 (1.2.2 ECG算法现状) +[14]")

# ================================================================
# [18] 程建华 — 高精度温度测量补偿算法
# ================================================================
# IDX=74: Section 2.1.3 体温检测原理 (already has [17])
p = doc.paragraphs[74]
if "[18]" not in p.text:
    add_citation(p, "[18]")
    results.append(f"P74 (2.1.3 体温检测原理) +[18]")

# IDX=186: Section 4.3.1 FIR低通滤波实现
p = doc.paragraphs[186]
if "[18]" not in p.text:
    add_citation(p, "[18]")
    results.append(f"P186 (4.3.1 FIR低通滤波) +[18]")

# IDX=188: Section 4.3.2 NTC校准 (already has [17])
p = doc.paragraphs[188]
if "[18]" not in p.text:
    add_citation(p, "[18]")
    results.append(f"P188 (4.3.2 NTC校准) +[18]")

# ================================================================
# [21] Yan H — QRS detection with adaptively regularized numerical differentiation
# ================================================================
# IDX=52: Section 1.2.2 ECG算法研究现状
p = doc.paragraphs[52]
if "[21]" not in p.text:
    add_citation(p, "[21]")
    results.append(f"P52 (1.2.2 ECG算法现状) +[21]")

# IDX=136: Section 4.1.2 差分阈值R波检测算法
p = doc.paragraphs[136]
if "[21]" not in p.text:
    add_citation(p, "[21]")
    results.append(f"P136 (4.1.2 差分阈值算法) +[21]")

# ================================================================
doc.save(SRC)
print(f"Done! Saved to: {SRC}")
print(f"\nAdded {len(results)} citations (all in RED font):")
for r in results:
    print(f"  {r}")
